"""Test cases from CSV/XLSX: one row = one chunk."""

from __future__ import annotations

import csv
import io
import re

from qabuddy.ingest.chunking import estimate_tokens, split_by_size
from qabuddy.ingest.context import ParseContext, tag
from qabuddy.ingest.models import Chunk, SourceFile, make_chunk_id
from qabuddy.ingest.normalize import clean_text, mask_secrets, read_text

FIELD_PATTERNS: list[tuple[str, re.Pattern]] = [
    ("id", re.compile(r"(^|\s)(t?id|tc ?id|test ?case ?id|scenario ?t?id|key)$")),
    ("title", re.compile(r"description|title|summary|scenario name|test ?case name")),
    ("precondition", re.compile(r"pre-?condition")),
    ("steps", re.compile(r"step")),
    ("expected", re.compile(r"expected")),
    ("actual", re.compile(r"actual")),
    ("status", re.compile(r"^status$|result status|execution status")),
    ("priority", re.compile(r"priority|severity")),
    ("automated", re.compile(r"automat")),
    ("executed_by", re.compile(r"qa name|executed by|tester")),
    ("comments", re.compile(r"comment|misc|notes")),
]


def _field_for(header: str) -> str | None:
    name = header.strip().lower()
    for field, pattern in FIELD_PATTERNS:
        if pattern.search(name):
            return field
    return None


def _read_rows(sf: SourceFile) -> list[list[str]]:
    if sf.abs_path.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(sf.abs_path, read_only=True, data_only=True)
        sheet = workbook.worksheets[0]
        return [["" if v is None else str(v) for v in row] for row in sheet.iter_rows(values_only=True)]
    return list(csv.reader(io.StringIO(read_text(sf.abs_path).lstrip("﻿"))))


def _merge_values(values: list[str]) -> str:
    kept: list[str] = []
    for value in values:
        value = value.strip()
        if value and value not in kept:
            kept.append(value)
    if len(kept) <= 1:
        return kept[0] if kept else ""
    return kept[0] + "".join(f"\n({i}) {v}" for i, v in enumerate(kept[1:], 2))


def _format_steps(steps: str) -> str:
    parts = [p.strip() for p in re.split(r"\s*\|\s*|\n", steps) if p.strip()]
    if not parts:
        return ""
    if all(re.match(r"^\d+[.)]", p) for p in parts):
        return "\n".join(parts)
    return "\n".join(f"{i}. {p}" for i, p in enumerate(parts, 1))


def _module(tc_id: str) -> str:
    return tc_id.split("-", 1)[0].title() if "-" in tc_id else ""


def parse_testcases(sf: SourceFile, ctx: ParseContext) -> list[Chunk]:
    rows = _read_rows(sf)
    if not rows:
        return []
    header, body = rows[0], rows[1:]
    fields = [_field_for(h) for h in header]
    if "id" not in fields:
        fields[0] = "id"
    settings = ctx.settings.chunking
    file_name = sf.rel_path.rsplit("/", 1)[-1]
    chunks: list[Chunk] = []
    seen_ids: set[str] = set()
    for row_number, row in enumerate(body, 1):
        values: dict[str, list[str]] = {}
        extras: list[tuple[str, str]] = []
        for name, field, value in zip(header, fields, row):
            if field:
                values.setdefault(field, []).append(value)
            elif value.strip():
                extras.append((name.strip(), value.strip()))
        record = {field: _merge_values(vals) for field, vals in values.items()}
        tc_id = record.get("id", "").strip() or f"ROW-{row_number}"
        locator = tc_id if tc_id not in seen_ids else f"{tc_id}@{row_number}"
        seen_ids.add(tc_id)

        meta = [f"Module: {_module(tc_id)}" if _module(tc_id) else ""]
        for label, key in (("Priority", "priority"), ("Automated", "automated"), ("Status", "status")):
            if record.get(key):
                meta.append(f"{label}: {record[key]}")
        lines = [f"Test case {tc_id} — {record.get('title', '')}".rstrip(" —"), " | ".join(m for m in meta if m)]
        if record.get("precondition"):
            lines.append(f"Precondition: {record['precondition']}")
        if record.get("steps"):
            lines.append("Steps:\n" + _format_steps(record["steps"]))
        if record.get("expected"):
            lines.append(f"Expected: {record['expected']}")
        if record.get("actual"):
            lines.append(f"Actual: {record['actual']}")
        if record.get("comments"):
            lines.append(f"Comments: {record['comments']}")
        if record.get("executed_by"):
            lines.append(f"Executed by: {record['executed_by']}")
        lines.extend(f"{name}: {value}" for name, value in extras)
        text = mask_secrets(clean_text("\n".join(line for line in lines if line)), kv_lines=False)

        pieces = [(text, 0, 0)] if estimate_tokens(text) <= settings.max_tokens else split_by_size(
            text, settings.target_tokens, settings.max_tokens, settings.prose_overlap_tokens)
        for index, (piece, _, _) in enumerate(pieces):
            chunk = Chunk(
                chunk_id=make_chunk_id(sf.rel_path, locator if index == 0 else f"{locator}#{index}"),
                source=sf.source,
                kind="testcase",
                path=sf.rel_path,
                citation=tc_id,
                text=piece,
                header=f"{ctx.source_label or 'Test cases'} › {file_name} › {tc_id}",
                location={"row": row_number},
                title=record.get("title", ""),
                parent_id=tc_id,
                module=_module(tc_id),
            )
            chunks.append(tag(chunk, ctx, [tc_id]))
    return chunks


def entities_from_testcases(chunks: list[Chunk]) -> set[str]:
    return {c.citation.lower() for c in chunks if c.kind == "testcase"}
